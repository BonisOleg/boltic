"""Читання / запис файлів каталогу: xlsx, csv, json, xml."""

from __future__ import annotations

import csv
import io
import json
import xml.etree.ElementTree as ET
from typing import Any
from xml.dom import minidom

from apps.catalog.io_schema import (
    FORMAT_ID,
    FORMAT_VERSION,
    FULL_DUMP_FIELDS,
    cell_to_export,
    detect_format,
    normalize_header,
)


SUPPORTED_EXTENSIONS = (".xlsx", ".csv", ".json", ".xml")


class CatalogIOError(Exception):
    pass


def extension_of(filename: str) -> str:
    name = (filename or "").lower().strip()
    for ext in SUPPORTED_EXTENSIONS:
        if name.endswith(ext):
            return ext
    return ""


def read_tabular_file(filename: str, content: bytes) -> tuple[str, list[dict[str, Any]]]:
    """
    Повертає (detected_format, rows_as_dicts).
    detected_format: 'goods' | 'full'
    """
    ext = extension_of(filename)
    if not ext:
        raise CatalogIOError(
            "Підтримуються лише файли: " + ", ".join(SUPPORTED_EXTENSIONS)
        )
    if ext == ".xlsx":
        headers, matrix = _read_xlsx(content)
    elif ext == ".csv":
        headers, matrix = _read_csv(content)
    elif ext == ".json":
        headers, matrix = _read_json(content)
    else:
        headers, matrix = _read_xml(content)

    fmt = detect_format(headers)
    if not fmt:
        raise CatalogIOError(
            "Невідома схема файлу. Очікується повний дамп (колонка article) "
            "або goods.xlsx (колонки «Код», «Ім'я товару»)."
        )
    rows: list[dict[str, Any]] = []
    for values in matrix:
        row: dict[str, Any] = {}
        for i, header in enumerate(headers):
            key = normalize_header(header) if fmt == "goods" else str(header).strip()
            if not key:
                continue
            row[key] = values[i] if i < len(values) else None
        rows.append(row)
    return fmt, rows


def write_file(fmt: str, rows: list[dict[str, Any]]) -> tuple[bytes, str, str]:
    """
    Повертає (content, content_type, filename).
    fmt: xlsx|csv|json|xml
    """
    fmt = fmt.lower().strip(".")
    if fmt == "xlsx":
        return _write_xlsx(rows), (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ), "catalog_full.xlsx"
    if fmt == "csv":
        return _write_csv(rows), "text/csv; charset=utf-8", "catalog_full.csv"
    if fmt == "json":
        return _write_json(rows), "application/json; charset=utf-8", "catalog_full.json"
    if fmt == "xml":
        return _write_xml(rows), "application/xml; charset=utf-8", "catalog_full.xml"
    raise CatalogIOError(f"Невідомий формат експорту: {fmt}")


def _read_xlsx(content: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise CatalogIOError("Потрібен openpyxl") from exc
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheet_name = "Список товарів" if "Список товарів" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise CatalogIOError("Порожній файл") from exc
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    matrix: list[list[Any]] = []
    for row in rows_iter:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        matrix.append(list(row))
    return headers, matrix


def _read_csv(content: bytes) -> tuple[list[str], list[list[Any]]]:
    text = content.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration as exc:
        raise CatalogIOError("Порожній CSV") from exc
    matrix = [list(row) for row in reader if any(str(c).strip() for c in row)]
    return headers, matrix


def _read_json(content: bytes) -> tuple[list[str], list[list[Any]]]:
    data = json.loads(content.decode("utf-8"))
    if isinstance(data, dict):
        items = data.get("items") or data.get("skus") or data.get("data")
    elif isinstance(data, list):
        items = data
    else:
        raise CatalogIOError("JSON має бути масивом або об'єктом з ключем items")
    if not isinstance(items, list) or not items:
        raise CatalogIOError("У JSON немає записів")
    if not isinstance(items[0], dict):
        raise CatalogIOError("Кожен запис JSON має бути об'єктом")
    headers = list(FULL_DUMP_FIELDS)
    # зберегти додаткові ключі з першого рядка
    for key in items[0].keys():
        if key not in headers:
            headers.append(key)
    matrix = [[item.get(h) for h in headers] for item in items]
    return headers, matrix


def _read_xml(content: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        root = ET.fromstring(content)
    except ET.ParseError as exc:
        raise CatalogIOError(f"Некоректний XML: {exc}") from exc
    nodes = list(root.findall("sku")) or list(root.findall("item"))
    if not nodes:
        raise CatalogIOError("У XML немає вузлів <sku> / <item>")
    headers = list(FULL_DUMP_FIELDS)
    for child in list(nodes[0]):
        if child.tag not in headers:
            headers.append(child.tag)
    matrix: list[list[Any]] = []
    for node in nodes:
        values = []
        for h in headers:
            el = node.find(h)
            values.append(el.text if el is not None else None)
        matrix.append(values)
    return headers, matrix


def _ordered_rows(rows: list[dict[str, Any]]) -> list[list[Any]]:
    return [[cell_to_export(row.get(f)) for f in FULL_DUMP_FIELDS] for row in rows]


def _write_xlsx(rows: list[dict[str, Any]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SKU"
    ws.append(list(FULL_DUMP_FIELDS))
    for line in _ordered_rows(rows):
        ws.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_csv(rows: list[dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=",", lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(FULL_DUMP_FIELDS)
    for line in _ordered_rows(rows):
        writer.writerow(["" if v is None else v for v in line])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def _write_json(rows: list[dict[str, Any]]) -> bytes:
    payload = {
        "format": FORMAT_ID,
        "version": FORMAT_VERSION,
        "items": [
            {f: cell_to_export(row.get(f)) for f in FULL_DUMP_FIELDS} for row in rows
        ],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _write_xml(rows: list[dict[str, Any]]) -> bytes:
    root = ET.Element(
        "catalog",
        {"format": FORMAT_ID, "version": str(FORMAT_VERSION)},
    )
    for row in rows:
        sku_el = ET.SubElement(root, "sku")
        for field in FULL_DUMP_FIELDS:
            child = ET.SubElement(sku_el, field)
            value = cell_to_export(row.get(field))
            if value is None or value == "":
                child.text = ""
            elif isinstance(value, bool):
                child.text = "true" if value else "false"
            else:
                child.text = str(value)
    rough = ET.tostring(root, encoding="utf-8")
    pretty = minidom.parseString(rough).toprettyxml(indent="  ", encoding="utf-8")
    return pretty
